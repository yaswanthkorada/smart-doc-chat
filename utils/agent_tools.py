"""
Agent Tools Module
Provides specialized tools for each agent to use
"""

import os
import json
from typing import List, Dict, Any
from pathlib import Path
from loguru import logger
from crewai.tools import BaseTool
from pydantic import BaseModel, Field
import PyPDF2
import docx
import pandas as pd
from pptx import Presentation
from youtube_transcript_api import YouTubeTranscriptApi
from bs4 import BeautifulSoup
import requests
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
from config import config

# ==================== INGESTION TOOLS ====================

class DocumentAnalyzerInput(BaseModel):
    """Input schema for document analyzer"""
    file_path: str = Field(..., description="Path to the document file")

class DocumentAnalyzerTool(BaseTool):
    name: str = "Document Analyzer"
    description: str = "Analyzes document type, structure, and recommends extraction strategy"
    
    def _run(self, file_path: str) -> str:
        """Analyze document and provide processing recommendations"""
        try:
            path = Path(file_path)
            file_type = path.suffix.lower().strip('.')
            file_size = path.stat().st_size
            
            analysis = {
                "file_name": path.name,
                "file_type": file_type,
                "file_size_mb": round(file_size / (1024 * 1024), 2),
                "recommendations": []
            }
            
            # Type-specific analysis
            if file_type == 'pdf':
                with open(file_path, 'rb') as f:
                    pdf = PyPDF2.PdfReader(f)
                    page_count = len(pdf.pages)
                    analysis["page_count"] = page_count
                    analysis["recommendations"].append(f"Use PyPDF2 for {page_count} pages")
                    analysis["recommended_chunk_size"] = 1000 if page_count > 50 else 800
                    
            elif file_type in ['docx', 'doc']:
                analysis["recommendations"].append("Use python-docx for text extraction")
                analysis["recommended_chunk_size"] = 800
                
            elif file_type in ['xlsx', 'xls', 'csv']:
                analysis["recommendations"].append("Use pandas for structured data")
                analysis["recommended_chunk_size"] = 500
                analysis["special_handling"] = "Preserve table structure"
                
            elif file_type in ['pptx', 'ppt']:
                prs = Presentation(file_path)
                slide_count = len(prs.slides)
                analysis["slide_count"] = slide_count
                analysis["recommendations"].append(f"Extract text from {slide_count} slides")
                analysis["recommended_chunk_size"] = 600
            
            return json.dumps(analysis, indent=2)
            
        except Exception as e:
            logger.error(f"Document analysis error: {e}")
            return f"Error analyzing document: {str(e)}"


class TextExtractorInput(BaseModel):
    """Input schema for text extractor"""
    file_path: str = Field(..., description="Path to the document file")
    strategy: str = Field(default="standard", description="Extraction strategy: standard, detailed, or quick")

class TextExtractorTool(BaseTool):
    name: str = "Text Extractor"
    description: str = "Extracts text from various document types (PDF, DOCX, PPTX, XLSX, TXT)"
    
    def _run(self, file_path: str, strategy: str = "standard") -> str:
        """Extract text based on file type"""
        try:
            path = Path(file_path)
            file_type = path.suffix.lower().strip('.')
            extracted_text = ""
            
            if file_type == 'pdf':
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page_num, page in enumerate(pdf_reader.pages):
                        text = page.extract_text()
                        extracted_text += f"\n--- Page {page_num + 1} ---\n{text}\n"
            
            elif file_type in ['docx', 'doc']:
                doc = docx.Document(file_path)
                extracted_text = '\n\n'.join([para.text for para in doc.paragraphs])
            
            elif file_type in ['pptx', 'ppt']:
                prs = Presentation(file_path)
                for slide_num, slide in enumerate(prs.slides):
                    extracted_text += f"\n--- Slide {slide_num + 1} ---\n"
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            extracted_text += shape.text + "\n"
            
            elif file_type in ['xlsx', 'xls']:
                df = pd.read_excel(file_path)
                extracted_text = f"Data Summary:\n{df.to_string()}\n\nStatistics:\n{df.describe().to_string()}"
            
            elif file_type == 'csv':
                df = pd.read_csv(file_path)
                extracted_text = f"Data Summary:\n{df.to_string()}\n\nStatistics:\n{df.describe().to_string()}"
            
            elif file_type == 'txt':
                with open(file_path, 'r', encoding='utf-8') as f:
                    extracted_text = f.read()
            
            else:
                return f"Unsupported file type: {file_type}"
            
            logger.info(f"Extracted {len(extracted_text)} characters from {path.name}")
            return extracted_text[:50000]  # Limit to 50k chars for agent processing
            
        except Exception as e:
            logger.error(f"Text extraction error: {e}")
            return f"Error extracting text: {str(e)}"


class DocumentChunkerInput(BaseModel):
    """Input schema for document chunker"""
    text: str = Field(..., description="Text to split into chunks")
    chunk_size: int = Field(default=1000, description="Size of each chunk")
    chunk_overlap: int = Field(default=200, description="Overlap between chunks")

class DocumentChunkerTool(BaseTool):
    name: str = "Document Chunker"
    description: str = "Splits text into optimal chunks for embedding"
    
    def _run(self, text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> str:
        """Split text into chunks"""
        try:
            text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                length_function=len
            )
            
            chunks = text_splitter.split_text(text)
            
            result = {
                "total_chunks": len(chunks),
                "avg_chunk_size": sum(len(c) for c in chunks) // len(chunks) if chunks else 0,
                "chunk_preview": chunks[0][:200] + "..." if chunks else "",
                "chunks": chunks[:100]  # Limit to 100 chunks for agent
            }
            
            logger.info(f"Created {len(chunks)} chunks")
            return json.dumps(result, indent=2)
            
        except Exception as e:
            logger.error(f"Chunking error: {e}")
            return f"Error chunking text: {str(e)}"


# ==================== RETRIEVAL TOOLS ====================

class VectorSearchInput(BaseModel):
    """Input schema for vector search"""
    query: str = Field(..., description="Search query")
    collection_name: str = Field(..., description="ChromaDB collection name")
    n_results: int = Field(default=5, description="Number of results to return")

class VectorSearchTool(BaseTool):
    name: str = "Vector Search"
    description: str = "Searches vector database for semantically similar content"
    
    def _run(self, query: str, collection_name: str, n_results: int = 5) -> str:
        """Search ChromaDB for relevant chunks"""
        try:
            from langchain_community.vectorstores import Chroma
            from langchain_huggingface import HuggingFaceEmbeddings
            
            # Initialize embeddings
            embeddings = HuggingFaceEmbeddings(
                model_name=config.HUGGINGFACE_MODEL,
                model_kwargs={'device': 'cpu'},
                encode_kwargs={'normalize_embeddings': True}
            )
            
            # Load collection
            persist_dir = os.path.join(config.CHROMA_PERSIST_DIR, collection_name)
            
            if not os.path.exists(persist_dir):
                return f"Collection not found: {collection_name}"
            
            vectorstore = Chroma(
                persist_directory=persist_dir,
                embedding_function=embeddings
            )
            
            # Search
            results = vectorstore.similarity_search_with_score(query, k=n_results)
            
            formatted_results = []
            for doc, score in results:
                formatted_results.append({
                    "content": doc.page_content[:500],
                    "relevance_score": round(float(score), 3),
                    "metadata": doc.metadata
                })
            
            return json.dumps(formatted_results, indent=2)
            
        except Exception as e:
            logger.error(f"Vector search error: {e}")
            return f"Error searching: {str(e)}"


class QueryAnalyzerInput(BaseModel):
    """Input schema for query analyzer"""
    query: str = Field(..., description="User's query to analyze")

class QueryAnalyzerTool(BaseTool):
    name: str = "Query Analyzer"
    description: str = "Analyzes user queries to understand intent and improve search"
    
    def _run(self, query: str) -> str:
        """Analyze query and suggest improvements"""
        try:
            analysis = {
                "original_query": query,
                "query_length": len(query),
                "query_type": self._detect_query_type(query),
                "key_terms": self._extract_key_terms(query),
                "suggested_reformulations": self._suggest_reformulations(query)
            }
            
            return json.dumps(analysis, indent=2)
            
        except Exception as e:
            logger.error(f"Query analysis error: {e}")
            return f"Error analyzing query: {str(e)}"
    
    def _detect_query_type(self, query: str) -> str:
        """Detect the type of query"""
        query_lower = query.lower()
        
        if any(word in query_lower for word in ['what', 'who', 'when', 'where', 'which']):
            return "factual"
        elif any(word in query_lower for word in ['how', 'why']):
            return "explanatory"
        elif any(word in query_lower for word in ['compare', 'difference', 'versus', 'vs']):
            return "comparative"
        elif any(word in query_lower for word in ['summarize', 'summary', 'overview']):
            return "summarization"
        else:
            return "general"
    
    def _extract_key_terms(self, query: str) -> List[str]:
        """Extract key terms from query"""
        # Simple keyword extraction (could use NER or more advanced methods)
        stop_words = {'what', 'is', 'the', 'a', 'an', 'how', 'why', 'when', 'where', 'who', 'which', 'are', 'was', 'were'}
        words = query.lower().split()
        key_terms = [w for w in words if w not in stop_words and len(w) > 3]
        return key_terms[:5]
    
    def _suggest_reformulations(self, query: str) -> List[str]:
        """Suggest alternative query formulations"""
        reformulations = []
        
        # Add more specific versions
        if "what is" in query.lower():
            reformulations.append(query.replace("What is", "Define"))
        
        # Add broader versions
        key_terms = self._extract_key_terms(query)
        if key_terms:
            reformulations.append(f"Information about {' '.join(key_terms)}")
        
        return reformulations[:3]


# ==================== GENERATION TOOLS ====================

class CitationFormatterInput(BaseModel):
    """Input schema for citation formatter"""
    text: str = Field(..., description="Text to format with citations")
    sources: str = Field(..., description="JSON string of source metadata")

class CitationFormatterTool(BaseTool):
    name: str = "Citation Formatter"
    description: str = "Formats responses with proper citations and source references"
    
    def _run(self, text: str, sources: str) -> str:
        """Add citations to text"""
        try:
            source_list = json.loads(sources) if isinstance(sources, str) else sources
            
            # Add source references at the end
            citation_text = text + "\n\n**Sources:**\n"
            
            for idx, source in enumerate(source_list, 1):
                metadata = source.get('metadata', {})
                filename = metadata.get('source', 'Unknown')
                page = metadata.get('page', 'N/A')
                
                citation_text += f"{idx}. {Path(filename).name}"
                if page != 'N/A':
                    citation_text += f" (Page {page})"
                citation_text += "\n"
            
            return citation_text
            
        except Exception as e:
            logger.error(f"Citation formatting error: {e}")
            return text  # Return original text if formatting fails


class ResponseValidatorInput(BaseModel):
    """Input schema for response validator"""
    response: str = Field(..., description="Generated response to validate")
    source_chunks: str = Field(..., description="Original source chunks as JSON")

class ResponseValidatorTool(BaseTool):
    name: str = "Response Validator"
    description: str = "Validates response accuracy against source material"
    
    def _run(self, response: str, source_chunks: str) -> str:
        """Validate response against sources"""
        try:
            chunks = json.loads(source_chunks) if isinstance(source_chunks, str) else source_chunks
            
            validation = {
                "response_length": len(response),
                "has_citations": "[" in response or "Source" in response,
                "potential_hallucinations": [],
                "validation_score": 0.0
            }
            
            # Check if response content appears in sources
            response_sentences = response.split('.')
            matched_sentences = 0
            
            for sentence in response_sentences:
                if len(sentence.strip()) < 10:
                    continue
                    
                found = False
                for chunk in chunks:
                    chunk_text = chunk.get('content', '')
                    # Check for semantic overlap (simplified)
                    if any(word in chunk_text.lower() for word in sentence.lower().split() if len(word) > 4):
                        found = True
                        matched_sentences += 1
                        break
                
                if not found and len(sentence.strip()) > 20:
                    validation["potential_hallucinations"].append(sentence.strip()[:100])
            
            # Calculate validation score
            if response_sentences:
                validation["validation_score"] = round(matched_sentences / len(response_sentences), 2)
            
            return json.dumps(validation, indent=2)
            
        except Exception as e:
            logger.error(f"Response validation error: {e}")
            return f"Error validating response: {str(e)}"


# Export all tools
def get_ingestion_tools():
    """Get tools for ingestion agent"""
    return [
        DocumentAnalyzerTool(),
        TextExtractorTool(),
        DocumentChunkerTool()
    ]

def get_retrieval_tools():
    """Get tools for retrieval agent"""
    return [
        QueryAnalyzerTool(),
        VectorSearchTool()
    ]

def get_generation_tools():
    """Get tools for generation agent"""
    return [
        CitationFormatterTool(),
        ResponseValidatorTool()
    ]
